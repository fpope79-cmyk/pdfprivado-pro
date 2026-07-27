from __future__ import annotations
import json, math, statistics, subprocess, os, shutil, tempfile
from pathlib import Path
from functools import lru_cache
from xml.sax.saxutils import escape

import fitz
import numpy as np
from PIL import Image, ImageFont
from skimage.metrics import structural_similarity as ssim
from docx import Document
from docx.oxml import parse_xml
from docx.shared import Pt
from docx.enum.section import WD_SECTION_START

EMU=12700
TWIP=20
ARIMO='/usr/share/fonts/truetype/croscore/Arimo-Regular.ttf'

@lru_cache(maxsize=1000)
def fnt(px): return ImageFont.truetype(ARIMO,max(1,int(round(px*10))))
def width(text,size):
    b=fnt(size).getbbox(str(text or '')); return max(.01,(b[2]-b[0])/10)
def clean(v): return ' '.join(str(v or '').split())
def clamp(v,a,b): return max(a,min(b,float(v)))

def effective_segments(line):
    seg=[s for s in line.get('segments',[]) if clean(s.get('text'))]
    if seg:return seg
    return [{'text':clean(line.get('text')),'x':line.get('x',0),'width':max(1,line.get('width',1))}]

def initial_size(line,page):
    raw=clamp(line.get('fontSize',9),1,96)
    layout=page.get('layout',[])
    sizes=sorted(float(e.get('fontSize',0)) for e in layout if float(e.get('fontSize',0))>0)
    med=sizes[len(sizes)//2] if sizes else raw
    L=len(clean(line.get('text')).replace(' ',''))
    mx=med*1.35 if L>=55 else med*1.55 if L>=28 else med*1.80 if L>=18 else math.inf
    raw=min(raw,mx)
    sparse=len(layout)<=32
    area=max(1,page['width']*page['height'])
    cov=sum(max(0,float(e.get('width',0)))*max(0,float(e.get('height',e.get('fontSize',0)))) for e in layout)/area
    return raw*(1.35/.82 if sparse and cov>=.24 else 1)

def display_size(line,page):
    raw=initial_size(line,page)
    seg=effective_segments(line)
    target=sum(max(0,float(s.get('width',0))) for s in seg) if len(seg)>1 else max(1,float(seg[0].get('width',line.get('width',1))))
    measured=sum(width(s['text'],raw) for s in seg)
    fit=clamp((target/max(1,measured))*.98,.58,1.08)
    sparse=len(page.get('layout',[]))<=32
    base=.78 if sparse else .86
    blend=.72 if sparse else .42
    return max(4.2,raw*(fit*blend+base*(1-blend)))

def segment_scale(text,target,size,conf):
    if len(clean(text).replace(' ',''))<3 or target<=0:return None
    if conf is not None and conf<70:return None
    m=width(text,size); ratio=target/m
    if ratio<.90 or ratio>1.10:return None
    pc=round(ratio*100); return pc if abs(pc-100)>=2 else None

def word_scale(text,target,size,conf):
    if not clean(text) or target<=0:return None
    if conf is not None and conf<55:return None
    m=width(text,size); ratio=target/m
    if ratio<.78 or ratio>1.22:return None
    pc=round(ratio*100); return pc if abs(pc-100)>=2 else None

def predict_v14_mae(line,page,size):
    runs=sorted([r for r in line.get('styledRuns',[]) if clean(r.get('text')) and 'x' in r and 'width' in r],key=lambda r:float(r['x']))
    if len(runs)<2:return 0,0
    pred=[];actual=[]
    for seg in effective_segments(line):
        x0=float(seg.get('x',line.get('x',0))); x1=x0+float(seg.get('width',1))
        sr=[r for r in runs if x0-.2 <= float(r['x'])+float(r['width'])/2 <= x1+.2]
        if not sr:continue
        full=' '.join(clean(r['text']) for r in sr)
        sc=segment_scale(full,float(seg.get('width',1)),size,float(line.get('confidence') or 0)) or 100
        prefix=''
        for i,r in enumerate(sr):
            if i:prefix+=' '
            pred.append(x0+width(prefix,size)*sc/100)
            actual.append(float(r['x']))
            prefix+=clean(r['text'])
    if len(pred)!=len(runs):return 0,0
    es=[abs(a-b) for a,b in zip(actual,pred)]
    return sum(es)/len(es),max(es)

def anchor_enabled(line,page,size):
    conf=float(line.get('confidence') or 0)
    runs=sorted([r for r in line.get('styledRuns',[]) if clean(r.get('text')) and 'x' in r and 'width' in r],key=lambda r:float(r['x']))
    if conf<75 or len(runs)<2 or len(runs)>28:return False
    if sum(len(clean(r['text']).replace(' ','')) for r in runs)<5:return False
    if any(float(runs[i]['x'])<=float(runs[i-1]['x'])+.35 for i in range(1,len(runs))):return False
    mae,mx=predict_v14_mae(line,page,size)
    return mae>=max(.75,size*.075) or mx>=max(1.8,size*.17)

def run_xml(text,font_size,bold=False,italic=False,color='000000',scale=None):
    scale_xml=f'<w:w w:val="{int(scale)}"/>' if scale else ''
    return f'''<w:r><w:rPr><w:rFonts w:ascii="Arial" w:cs="Arial" w:eastAsia="Arial" w:hAnsi="Arial"/><w:b w:val="{'true' if bold else 'false'}"/><w:bCs w:val="{'true' if bold else 'false'}"/><w:i w:val="{'true' if italic else 'false'}"/><w:iCs w:val="{'true' if italic else 'false'}"/><w:noProof/><w:color w:val="{color}"/><w:sz w:val="{max(1,round(font_size*2))}"/><w:szCs w:val="{max(1,round(font_size*2))}"/>{scale_xml}</w:rPr><w:t xml:space="preserve">{escape(str(text))}</w:t></w:r>'''

def textbox_xml(line,page,idx,mode):
    fs=initial_size(line,page); ds=display_size(line,page)
    x=float(line.get('x',0)); pageH=float(page['height'])
    # v14 old geometry: union bottom at base factor .82
    if mode in {'v14','v15-anchor'}:
        top=float(line.get('ocrTop',0))+max(fs,float(line.get('height',fs)))-fs*.82
    else:
        # robust baseline from median run bottoms; fallback union
        runs=[r for r in line.get('styledRuns',[]) if 'y' in r]
        hint=statistics.median(pageH-float(r['y']) for r in runs) if runs else float(line.get('ocrBottom',0))
        baseline_factor={'v16':.82,'v16-b94':.94,'v16-b106':1.06,'v16-b118':1.18,'v16-b124':1.24,'v16-b130':1.30}.get(mode,.82)
        top=hint-ds*.82-fs*(baseline_factor-.82)
    w=max(12,float(line.get('width',len(clean(line.get('text')))*fs*.45)))
    h=max(fs*1.30,ds*1.35,6)
    color=str(line.get('color') or '000000').replace('#','').upper(); color=color if len(color)==6 else '000000'
    bold=bool(line.get('bold')); italic=bool(line.get('italic')); conf=float(line.get('confidence') or 0)
    tabs=[]; runs_xml=[]
    if mode != 'v14' and anchor_enabled(line,page,ds):
        wr=sorted([r for r in line.get('styledRuns',[]) if clean(r.get('text')) and 'x' in r and 'width' in r],key=lambda r:float(r['x']))
        origin=float(wr[0]['x'])
        for i,r in enumerate(wr):
            if i:
                tabs.append(round((float(r['x'])-origin)*20))
                runs_xml.append(run_xml('\t',ds,bold,italic,color,None))
            sc=word_scale(clean(r['text']),float(r['width']),ds,float(r.get('confidence') or conf))
            runs_xml.append(run_xml(clean(r['text']),ds,bold,italic,color,sc))
    else:
        seg=effective_segments(line); origin=float(seg[0].get('x',x))
        for i,s in enumerate(seg):
            if i:
                tabs.append(round((float(s.get('x',origin))-origin)*20))
                runs_xml.append(run_xml('\t',ds,bold,italic,color,None))
            sc=segment_scale(clean(s['text']),float(s.get('width',1)),ds,conf)
            runs_xml.append(run_xml(clean(s['text']),ds,bold,italic,color,sc))
    tabs_xml='<w:tabs>'+''.join(f'<w:tab w:val="left" w:pos="{max(1,t)}"/>' for t in tabs)+'</w:tabs>' if tabs else ''
    # always left aligned for harness
    return f'''<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><w:pPr><w:spacing w:after="0" w:before="0" w:line="1" w:lineRule="exact"/></w:pPr><w:r><w:rPr><w:noProof/></w:rPr><w:drawing><wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="{251658240+idx}" behindDoc="0" locked="0" layoutInCell="0" allowOverlap="1"><wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="page"><wp:posOffset>{round(x*EMU)}</wp:posOffset></wp:positionH><wp:positionV relativeFrom="page"><wp:posOffset>{round(top*EMU)}</wp:posOffset></wp:positionV><wp:extent cx="{round(w*EMU)}" cy="{round(h*EMU)}"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/><wp:docPr id="{10000+idx}" name="Texto {idx}"/><wp:cNvGraphicFramePr><a:graphicFrameLocks noChangeAspect="0"/></wp:cNvGraphicFramePr><a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:wsp><wps:cNvSpPr txBox="1"/><wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{round(w*EMU)}" cy="{round(h*EMU)}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/><a:ln><a:noFill/></a:ln></wps:spPr><wps:txbx><w:txbxContent><w:p><w:pPr>{tabs_xml}<w:bidi w:val="false"/><w:spacing w:after="0" w:before="0" w:line="{max(1,round(max(ds*1.04,4.4)*20))}" w:lineRule="exact"/><w:jc w:val="left"/></w:pPr>{''.join(runs_xml)}</w:p></w:txbxContent></wps:txbx><wps:bodyPr rot="0" vert="horz" wrap="none" lIns="0" tIns="0" rIns="0" bIns="0" anchor="t" anchorCtr="0" upright="1"><a:noAutofit/></wps:bodyPr></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing></w:r></w:p>'''

def bg_xml(rid,page,idx):
    w=round(page['width']*EMU);h=round(page['height']*EMU)
    return f'''<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture"><w:pPr><w:spacing w:after="0" w:before="0" w:line="1" w:lineRule="exact"/></w:pPr><w:r><w:drawing><wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" allowOverlap="1" behindDoc="1" locked="0" layoutInCell="0" relativeHeight="1"><wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH><wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV><wp:extent cx="{w}" cy="{h}"/><wp:effectExtent t="0" r="0" b="0" l="0"/><wp:wrapNone/><wp:docPr id="{idx}" name="background"/><wp:cNvGraphicFramePr><a:graphicFrameLocks noChangeAspect="1"/></wp:cNvGraphicFramePr><a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture"><pic:pic><pic:nvPicPr><pic:cNvPr id="0" name="background"/><pic:cNvPicPr><a:picLocks noChangeAspect="1"/></pic:cNvPicPr></pic:nvPicPr><pic:blipFill><a:blip r:embed="{rid}"/><a:srcRect/><a:stretch><a:fillRect/></a:stretch></pic:blipFill><pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{w}" cy="{h}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr></pic:pic></a:graphicData></a:graphic></wp:anchor></w:drawing></w:r></w:p>'''

def make_doc(page,root,mode,out):
    doc=Document()
    sec=doc.sections[0];sec.page_width=Pt(page['width']);sec.page_height=Pt(page['height']);sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Pt(0);sec.header_distance=sec.footer_distance=Pt(0)
    # Un párrafo de flujo mínimo mantiene estable la página de Writer.
    p=doc.add_paragraph();p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(.05)
    bg=root/page['graphicsImage']
    rid,_=doc.part.get_or_add_image(str(bg))
    body=doc._element.body
    sect=body.sectPr
    body.insert(body.index(sect),parse_xml(bg_xml(rid,page,1)))
    for i,line in enumerate(page['layout'],1):
        body.insert(body.index(sect),parse_xml(textbox_xml(line,page,i,mode)))
    doc.save(out)

def render_docx(docx,outdir):
    subprocess.run(['libreoffice','--headless','--convert-to','pdf','--outdir',str(outdir),str(docx)],check=True,stdout=subprocess.PIPE,stderr=subprocess.PIPE)
    pdf=Path(outdir)/(Path(docx).stem+'.pdf')
    d=fitz.open(pdf);page=d[0];pix=page.get_pixmap(matrix=fitz.Matrix(2,2),alpha=False);arr=np.frombuffer(pix.samples,dtype=np.uint8).reshape(pix.height,pix.width,pix.n)[:,:,:3].copy();d.close();return arr

def compare(a,b):
    h=min(a.shape[0],b.shape[0]);w=min(a.shape[1],b.shape[1]);a=a[:h,:w];b=b[:h,:w]
    aa=np.dot(a[...,:3],[.299,.587,.114]);bb=np.dot(b[...,:3],[.299,.587,.114])
    return ssim(aa,bb,data_range=255),float(np.mean(np.abs(aa-bb)))

def main():
    root=Path('/mnt/data/pdfprivado-v14-bank/190604');data=json.load(open(root/'document.json'));page=data['pages'][0]
    work=Path('/mnt/data/docx-harness');shutil.rmtree(work,ignore_errors=True);work.mkdir()
    src=np.asarray(Image.open(root/page['sourceImage']).convert('RGB'))
    for mode in ['v14','v15-anchor','v16']:
        d=work/f'{mode}.docx';make_doc(page,root,mode,d);arr=render_docx(d,work);print(mode,arr.shape,src.shape,compare(src,arr))
if __name__=='__main__':main()

# --- multi-page harness -------------------------------------------------------
def make_doc_pages(pages, root, mode, out):
    doc = Document()
    first = pages[0]
    sec = doc.sections[0]
    sec.page_width = Pt(first['width']); sec.page_height = Pt(first['height'])
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Pt(0)
    sec.header_distance = sec.footer_distance = Pt(0)
    # tiny first paragraph
    p = doc.add_paragraph(); p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(.05)
    body=doc._element.body; sect=body.sectPr
    global_id=1
    for page_index,page in enumerate(pages):
        if page_index:
            pb=doc.add_paragraph(); pb.paragraph_format.space_before=Pt(0);pb.paragraph_format.space_after=Pt(0)
            run=pb.add_run(); run.add_break(__import__('docx').enum.text.WD_BREAK.PAGE)
        bg_path=root/page['graphicsImage']
        if not bg_path.exists():
            continue
        rid,_=doc.part.get_or_add_image(str(bg_path))
        global_id += 1
        body.insert(body.index(sect),parse_xml(bg_xml(rid,page,global_id)))
        for line in page.get('layout',[]):
            global_id += 1
            body.insert(body.index(sect),parse_xml(textbox_xml(line,page,global_id,mode)))
    doc.save(out)

def render_docx_pages(docx,outdir):
    subprocess.run(['libreoffice','--headless','--convert-to','pdf','--outdir',str(outdir),str(docx)],check=True,stdout=subprocess.PIPE,stderr=subprocess.PIPE)
    pdf=Path(outdir)/(Path(docx).stem+'.pdf')
    d=fitz.open(pdf); arrays=[]
    for page in d:
        pix=page.get_pixmap(matrix=fitz.Matrix(2,2),alpha=False)
        arrays.append(np.frombuffer(pix.samples,dtype=np.uint8).reshape(pix.height,pix.width,pix.n)[:,:,:3].copy())
    d.close();return arrays

def benchmark_directory(name, modes=('v14','v15-anchor','v16')):
    root=Path('/mnt/data/pdfprivado-v14-bank')/name
    candidate=root/'document.json'; geom=root/'document-geometry.json'; path=candidate if candidate.exists() else geom
    data=json.load(open(path)); pages=[p for p in data['pages'] if p.get('source')=='ocr' and (root/p.get('graphicsImage','')).exists()]
    work=Path('/mnt/data/docx-harness')/name;shutil.rmtree(work,ignore_errors=True);work.mkdir(parents=True)
    result={}
    for mode in modes:
        docx=work/f'{mode}.docx';make_doc_pages(pages,root,mode,docx);rendered=render_docx_pages(docx,work)
        vals=[]
        for i,page in enumerate(pages):
            if i>=len(rendered):break
            src=np.asarray(Image.open(root/page['sourceImage']).convert('RGB'))
            score,mae=compare(src,rendered[i]);vals.append({'page':page['pageNumber'],'ssim':float(score),'mae':mae})
        result[mode]=vals
    return result
